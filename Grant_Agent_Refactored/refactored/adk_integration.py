"""
Refactored Grant Agent - ADK Integration Layer
Clean integration with Google ADK framework
"""

import logging
import json
from typing import Dict, List, Optional, Any, Union
from datetime import datetime

from google.adk import Agent
from google.adk.core import Artifact

from .models import (
    GrantWorkflowData, WorkflowResult, DocumentData, VersionedWorkflowData,
    WorkflowStatus, DocumentType
)
from .services import DocumentService, ExtractionService, LanguageAnalysisService
from .workflow_engine import (
    WorkflowEngine, DocumentProcessingStage, DataExtractionStage,
    LanguageAnalysisStage, NarrativeGenerationStage, DocumentCreationStage,
    ConflictResolver
)

logger = logging.getLogger(__name__)


class LLMService:
    """Wrapper for ADK LLM interactions"""
    
    def __init__(self, adk_agent):
        self.agent = adk_agent
        self.logger = logging.getLogger(__name__)
    
    def generate_text(self, prompt: str) -> str:
        """Generate text using the ADK agent's LLM"""
        try:
            # Use agent's built-in LLM generation
            response = self.agent.llm.generate(prompt)
            return response.text if hasattr(response, 'text') else str(response)
        except Exception as e:
            self.logger.error(f"LLM generation failed: {e}")
            raise


class DocumentGenerator:
    """Generate final grant documents"""
    
    def __init__(self):
        self.logger = logging.getLogger(__name__)
    
    def create_grant_document(self, org_info, funder_info, narrative: str) -> bytes:
        """Create final DOCX document"""
        try:
            from docx import Document
            from io import BytesIO
            
            # Create new document
            doc = Document()
            
            # Add title
            title = doc.add_heading(f'Grant Proposal to {funder_info.name}', 0)
            
            # Add organization section
            doc.add_heading('Organization Information', level=1)
            doc.add_paragraph(f'Organization: {org_info.name}')
            doc.add_paragraph(f'Mission: {org_info.mission}')
            if org_info.location:
                doc.add_paragraph(f'Location: {org_info.location}')
            
            # Add grant request section
            doc.add_heading('Grant Request', level=1)
            doc.add_paragraph(f'Requested Amount: ${funder_info.grant_amount:,}')
            doc.add_paragraph(f'Funding Type: {funder_info.funding_type}')
            
            # Add narrative
            doc.add_heading('Project Narrative', level=1)
            doc.add_paragraph(narrative)
            
            # Add closing
            doc.add_heading('Conclusion', level=1)
            doc.add_paragraph(
                f'We respectfully request ${funder_info.grant_amount:,} from {funder_info.name} '
                f'to support our mission and create meaningful impact in our community. '
                f'Thank you for considering our proposal.'
            )
            
            # Save to bytes
            doc_buffer = BytesIO()
            doc.save(doc_buffer)
            doc_buffer.seek(0)
            
            self.logger.info("Created grant document successfully")
            return doc_buffer.getvalue()
            
        except Exception as e:
            self.logger.error(f"Document generation failed: {e}")
            raise


class ArtifactService:
    """Manage artifacts for the refactored system"""
    
    def __init__(self, adk_agent):
        self.agent = adk_agent
        self.logger = logging.getLogger(__name__)
    
    def store(self, artifact_id: str, data: bytes, mime_type: str) -> None:
        """Store data as artifact"""
        try:
            artifact = Artifact(
                id=artifact_id,
                content=data,
                mime_type=mime_type
            )
            # Store in agent's artifact system
            self.agent.artifacts[artifact_id] = artifact
            self.logger.info(f"Stored artifact: {artifact_id}")
        except Exception as e:
            self.logger.warning(f"Failed to store artifact {artifact_id}: {e}")
    
    def retrieve(self, artifact_id: str) -> Optional[bytes]:
        """Retrieve artifact data"""
        try:
            artifact = self.agent.artifacts.get(artifact_id)
            if artifact:
                return artifact.content
            return None
        except Exception as e:
            self.logger.warning(f"Failed to retrieve artifact {artifact_id}: {e}")
            return None


class SessionManager:
    """Manage workflow state in session"""
    
    def __init__(self, adk_agent):
        self.agent = adk_agent
        self.logger = logging.getLogger(__name__)
        self.STATE_KEY = "grant_workflow_data"
        self.VERSION_KEY = "grant_workflow_versions"
    
    def load_workflow_data(self) -> VersionedWorkflowData:
        """Load workflow data from session"""
        try:
            # Load current data
            current_data = self.agent.session_state.get(self.STATE_KEY)
            if current_data:
                # Deserialize from dict
                workflow_data = self._dict_to_workflow_data(current_data)
            else:
                workflow_data = GrantWorkflowData()
            
            # Load versions
            versions_data = self.agent.session_state.get(self.VERSION_KEY, [])
            
            versioned_data = VersionedWorkflowData(current=workflow_data)
            # Note: version loading would need more complex serialization
            
            self.logger.info(f"Loaded workflow data: {workflow_data.status.value}")
            return versioned_data
            
        except Exception as e:
            self.logger.error(f"Failed to load workflow data: {e}")
            return VersionedWorkflowData(current=GrantWorkflowData())
    
    def save_workflow_data(self, versioned_data: VersionedWorkflowData) -> None:
        """Save workflow data to session"""
        try:
            # Convert to serializable dict
            current_dict = self._workflow_data_to_dict(versioned_data.current)
            
            # Save current data
            self.agent.session_state[self.STATE_KEY] = current_dict
            
            # Save basic version info (simplified)
            version_info = [{
                "id": v.id,
                "created_at": v.created_at.isoformat(),
                "reason": v.reason
            } for v in versioned_data.versions]
            self.agent.session_state[self.VERSION_KEY] = version_info
            
            self.logger.info(f"Saved workflow data: {versioned_data.current.status.value}")
            
        except Exception as e:
            self.logger.error(f"Failed to save workflow data: {e}")
    
    def _workflow_data_to_dict(self, data: GrantWorkflowData) -> Dict[str, Any]:
        """Convert workflow data to serializable dict"""
        return {
            "documents": {
                k: {
                    "content": v.content,
                    "filename": v.filename,
                    "mime_type": v.mime_type,
                    "uploaded_at": v.uploaded_at.isoformat(),
                    "content_hash": v.content_hash,
                    "document_type": v.document_type.value
                } for k, v in data.documents.items()
            },
            "organization": self._org_to_dict(data.organization) if data.organization else None,
            "funder": self._funder_to_dict(data.funder) if data.funder else None,
            "language_analysis": self._analysis_to_dict(data.language_analysis) if data.language_analysis else None,
            "narrative": data.narrative,
            "status": data.status.value,
            "errors": data.errors,
            "created_at": data.created_at.isoformat(),
            "last_modified": data.last_modified.isoformat()
        }
    
    def _dict_to_workflow_data(self, data_dict: Dict[str, Any]) -> GrantWorkflowData:
        """Convert dict back to workflow data"""
        from .models import OrganizationInfo, FunderInfo, LanguageAnalysis, ContactInfo, FinancialInfo
        
        data = GrantWorkflowData()
        
        # Restore documents
        for k, v in data_dict.get("documents", {}).items():
            doc = DocumentData(
                content=v["content"],
                filename=v["filename"],
                mime_type=v["mime_type"],
                uploaded_at=datetime.fromisoformat(v["uploaded_at"]),
                content_hash=v["content_hash"],
                document_type=DocumentType(v["document_type"])
            )
            data.documents[k] = doc
        
        # Restore organization
        if data_dict.get("organization"):
            data.organization = self._dict_to_org(data_dict["organization"])
        
        # Restore funder
        if data_dict.get("funder"):
            data.funder = self._dict_to_funder(data_dict["funder"])
        
        # Restore analysis
        if data_dict.get("language_analysis"):
            data.language_analysis = self._dict_to_analysis(data_dict["language_analysis"])
        
        # Restore other fields
        data.narrative = data_dict.get("narrative")
        data.status = WorkflowStatus(data_dict.get("status", "initialized"))
        data.errors = data_dict.get("errors", [])
        data.created_at = datetime.fromisoformat(data_dict["created_at"])
        data.last_modified = datetime.fromisoformat(data_dict["last_modified"])
        
        return data
    
    def _org_to_dict(self, org) -> Dict[str, Any]:
        """Convert organization to dict"""
        return {
            "name": org.name,
            "mission": org.mission,
            "background": org.background,
            "location": org.location,
            "contact_info": {
                "name": org.contact_info.name,
                "title": org.contact_info.title,
                "email": org.contact_info.email,
                "phone": org.contact_info.phone
            },
            "financials": {
                "annual_budget": org.financials.annual_budget,
                "funding_sources": org.financials.funding_sources,
                "revenue_streams": org.financials.revenue_streams
            },
            "leadership": org.leadership,
            "achievements": org.achievements,
            "target_population": org.target_population
        }
    
    def _dict_to_org(self, data: Dict[str, Any]):
        """Convert dict to organization"""
        from .models import OrganizationInfo, ContactInfo, FinancialInfo
        
        contact = ContactInfo(
            name=data["contact_info"]["name"],
            title=data["contact_info"]["title"],
            email=data["contact_info"]["email"],
            phone=data["contact_info"]["phone"]
        )
        
        financials = FinancialInfo(
            annual_budget=data["financials"]["annual_budget"],
            funding_sources=data["financials"]["funding_sources"],
            revenue_streams=data["financials"]["revenue_streams"]
        )
        
        return OrganizationInfo(
            name=data["name"],
            mission=data["mission"],
            background=data["background"],
            location=data["location"],
            contact_info=contact,
            financials=financials,
            leadership=data["leadership"],
            achievements=data["achievements"],
            target_population=data["target_population"]
        )
    
    def _funder_to_dict(self, funder) -> Dict[str, Any]:
        """Convert funder to dict"""
        return {
            "name": funder.name,
            "grant_amount": funder.grant_amount,
            "eligibility_criteria": funder.eligibility_criteria,
            "funding_priorities": funder.funding_priorities,
            "requirements": funder.requirements,
            "funding_type": funder.funding_type,
            "deadline": funder.deadline.isoformat() if funder.deadline else None
        }
    
    def _dict_to_funder(self, data: Dict[str, Any]):
        """Convert dict to funder"""
        from .models import FunderInfo
        
        return FunderInfo(
            name=data["name"],
            grant_amount=data["grant_amount"],
            eligibility_criteria=data["eligibility_criteria"],
            funding_priorities=data["funding_priorities"],
            requirements=data["requirements"],
            funding_type=data["funding_type"],
            deadline=datetime.fromisoformat(data["deadline"]) if data["deadline"] else None
        )
    
    def _analysis_to_dict(self, analysis) -> Dict[str, Any]:
        """Convert analysis to dict"""
        return {
            "alignment_score": analysis.alignment_score,
            "matching_themes": analysis.matching_themes,
            "recommendations": analysis.recommendations,
            "risk_areas": analysis.risk_areas,
            "generated_at": analysis.generated_at.isoformat()
        }
    
    def _dict_to_analysis(self, data: Dict[str, Any]):
        """Convert dict to analysis"""
        from .models import LanguageAnalysis
        
        return LanguageAnalysis(
            alignment_score=data["alignment_score"],
            matching_themes=data["matching_themes"],
            recommendations=data["recommendations"],
            risk_areas=data["risk_areas"],
            generated_at=datetime.fromisoformat(data["generated_at"])
        )


class RefactoredGrantAgent(Agent):
    """Main Grant Agent using refactored architecture"""
    
    def __init__(self):
        super().__init__(name="GrantWorkflowAgent")
        self.logger = logging.getLogger(__name__)
        
        # Initialize services
        self.artifact_service = ArtifactService(self)
        self.session_manager = SessionManager(self)
        self.llm_service = LLMService(self)
        self.document_generator = DocumentGenerator()
        
        # Initialize core services with dependency injection
        self.document_service = DocumentService(self.artifact_service)
        self.extraction_service = ExtractionService()
        self.analysis_service = LanguageAnalysisService()
        
        # Initialize workflow engine
        self.workflow_engine = WorkflowEngine()
        self.conflict_resolver = ConflictResolver()
        
        # Register workflow stages
        self._setup_workflow_stages()
        
        # Setup callbacks
        self._setup_callbacks()
        
        self.logger.info("RefactoredGrantAgent initialized")
    
    def _setup_workflow_stages(self):
        """Register all workflow stages"""
        self.workflow_engine.register_stage(
            DocumentProcessingStage(self.document_service)
        )
        self.workflow_engine.register_stage(
            DataExtractionStage(self.extraction_service)
        )
        self.workflow_engine.register_stage(
            LanguageAnalysisStage(self.analysis_service)
        )
        self.workflow_engine.register_stage(
            NarrativeGenerationStage(self.llm_service)
        )
        self.workflow_engine.register_stage(
            DocumentCreationStage(self.document_generator)
        )
    
    def _setup_callbacks(self):
        """Setup unified callback system"""
        
        @self.before_model_callback
        def process_uploads(ctx):
            """Process file uploads before model interaction"""
            file_uploads = []
            
            # Extract file uploads from context
            if hasattr(ctx, 'files') and ctx.files:
                for file_obj in ctx.files:
                    try:
                        file_data = file_obj.read()
                        filename = getattr(file_obj, 'name', 'unknown.docx')
                        mime_type = getattr(file_obj, 'content_type', 
                                          'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
                        
                        file_uploads.append((file_data, filename, mime_type))
                        self.logger.info(f"Prepared file upload: {filename}")
                        
                    except Exception as e:
                        self.logger.error(f"Error processing file upload: {e}")
            
            # Store in context for workflow
            ctx.file_uploads = file_uploads
            
            return ctx
        
        @self.after_model_callback
        def execute_workflow(ctx):
            """Execute workflow after model interaction"""
            try:
                # Load current workflow state
                versioned_data = self.session_manager.load_workflow_data()
                
                # Build execution context
                execution_context = {
                    'file_uploads': getattr(ctx, 'file_uploads', []),
                    'user_input': str(ctx.user_message) if hasattr(ctx, 'user_message') else "",
                    'session_id': getattr(ctx, 'session_id', 'default')
                }
                
                # Execute workflow
                result = self.workflow_engine.execute_workflow(
                    versioned_data.current, 
                    execution_context
                )
                
                # Handle result
                if result.success:
                    # Update versioned data
                    versioned_data.current = result.data
                    
                    # Save state
                    self.session_manager.save_workflow_data(versioned_data)
                    
                    # Create response
                    response = self._create_workflow_response(result)
                    ctx.response = response
                    
                else:
                    # Handle workflow failure
                    error_response = f"Workflow failed at stage '{result.stage_completed}'"
                    if result.data.errors:
                        error_response += f"\nErrors: {'; '.join(result.data.errors[-3:])}"
                    
                    ctx.response = error_response
                
                self.logger.info(f"Workflow execution: {result.success}, Stage: {result.stage_completed}")
                
            except Exception as e:
                self.logger.error(f"Workflow execution failed: {e}")
                ctx.response = f"Internal workflow error: {e}"
            
            return ctx
    
    def _create_workflow_response(self, result: WorkflowResult) -> str:
        """Create user-friendly response based on workflow result"""
        data = result.data
        stage = result.stage_completed
        
        if result.has_conflicts():
            return self._create_conflict_response(result.conflicts)
        
        # Status-based responses
        if data.status == WorkflowStatus.DOCUMENTS_UPLOADED:
            doc_count = len(data.documents)
            doc_types = [doc.document_type.value for doc in data.documents.values()]
            return f"✅ Processed {doc_count} document(s): {', '.join(set(doc_types))}. Starting data extraction..."
        
        elif data.status == WorkflowStatus.EXTRACTION_COMPLETE:
            org_name = data.organization.name if data.organization else "Unknown"
            funder_name = data.funder.name if data.funder else "Unknown"
            return f"✅ Extracted information:\n• Organization: {org_name}\n• Funder: {funder_name}\nAnalyzing alignment..."
        
        elif data.status == WorkflowStatus.ANALYSIS_COMPLETE:
            score = data.language_analysis.alignment_score
            themes = ', '.join(data.language_analysis.matching_themes)
            return f"✅ Language analysis complete!\n• Alignment Score: {score:.1%}\n• Matching Themes: {themes}\nGenerating narrative..."
        
        elif data.status == WorkflowStatus.NARRATIVE_GENERATED:
            word_count = len(data.narrative.split()) if data.narrative else 0
            return f"✅ Generated {word_count}-word grant narrative! Creating final document..."
        
        elif data.status == WorkflowStatus.COMPLETED:
            return f"🎉 Grant proposal complete!\n• Organization: {data.organization.name}\n• Funder: {data.funder.name}\n• Amount: ${data.funder.grant_amount:,}\n\nYour grant proposal is ready for review!"
        
        else:
            return f"Workflow in progress... (Stage: {stage})"
    
    def _create_conflict_response(self, conflicts) -> str:
        """Create response for handling conflicts"""
        response = "⚠️ Found conflicts that need your input:\n\n"
        
        for i, conflict in enumerate(conflicts[:3], 1):  # Show first 3 conflicts
            response += f"{i}. {conflict.field_name}:\n"
            response += f"   Existing: {conflict.existing_value}\n"
            response += f"   New: {conflict.new_value}\n\n"
        
        response += "Please specify how to resolve these conflicts or I'll use the existing values."
        return response
    
    def get_status(self) -> Dict[str, Any]:
        """Get current workflow status"""
        try:
            versioned_data = self.session_manager.load_workflow_data()
            return self.workflow_engine.get_workflow_status(versioned_data.current)
        except Exception as e:
            self.logger.error(f"Error getting status: {e}")
            return {"error": str(e)}


# Factory function for creating the agent
def create_grant_agent() -> RefactoredGrantAgent:
    """Factory function to create and configure the grant agent"""
    agent = RefactoredGrantAgent()
    return agent 