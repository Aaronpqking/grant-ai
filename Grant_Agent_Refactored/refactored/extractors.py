"""
Text extraction utilities for document processing
"""

import logging
from typing import Optional

logger = logging.getLogger(__name__)


def extract_text_from_docx(file_data: bytes) -> str:
    """Extract text from DOCX file data"""
    try:
        from docx import Document
        from io import BytesIO
        
        # Create document from bytes
        doc_stream = BytesIO(file_data)
        doc = Document(doc_stream)
        
        # Extract all paragraph text
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        
        # Extract table text
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)
        
        extracted_text = '\n'.join(text_parts)
        logger.info(f"Extracted {len(extracted_text)} characters from DOCX")
        return extracted_text
        
    except Exception as e:
        logger.error(f"Error extracting DOCX text: {e}")
        return f"[Error extracting DOCX: {e}]"


def extract_text_from_pdf(file_data: bytes) -> str:
    """Extract text from PDF file data"""
    try:
        import PyPDF2
        from io import BytesIO
        
        # Create PDF reader from bytes
        pdf_stream = BytesIO(file_data)
        pdf_reader = PyPDF2.PdfReader(pdf_stream)
        
        # Extract text from all pages
        text_parts = []
        for page in pdf_reader.pages:
            page_text = page.extract_text()
            if page_text.strip():
                text_parts.append(page_text)
        
        extracted_text = '\n'.join(text_parts)
        logger.info(f"Extracted {len(extracted_text)} characters from PDF")
        return extracted_text
        
    except Exception as e:
        logger.error(f"Error extracting PDF text: {e}")
        return f"[Error extracting PDF: {e}]"


def extract_text_from_plain(file_data: bytes, encoding: str = 'utf-8') -> str:
    """Extract text from plain text file"""
    try:
        text = file_data.decode(encoding)
        logger.info(f"Extracted {len(text)} characters from plain text")
        return text
    except UnicodeDecodeError:
        try:
            # Try latin-1 encoding as fallback
            text = file_data.decode('latin-1')
            logger.info(f"Extracted {len(text)} characters from plain text (latin-1)")
            return text
        except Exception as e:
            logger.error(f"Error extracting plain text: {e}")
            return f"[Error extracting text: {e}]"
    except Exception as e:
        logger.error(f"Error extracting plain text: {e}")
        return f"[Error extracting text: {e}]" 