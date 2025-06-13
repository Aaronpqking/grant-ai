#!/usr/bin/env python3
"""
Create realistic sample documents for grant building demonstration
"""

from docx import Document
from pathlib import Path

def create_organization_document():
    """Create a realistic organization document"""
    doc = Document()
    
    doc.add_heading('Freedom Equity Inc - Organization Information', 0)
    
    doc.add_heading('Organization Overview', level=1)
    
    # Organization Name
    p = doc.add_paragraph()
    p.add_run('Organization Name: ').bold = True
    p.add_run('Freedom Equity Inc')
    
    # Mission Statement
    p = doc.add_paragraph()
    p.add_run('Mission Statement: ').bold = True
    p.add_run('To provide affordable housing solutions and small business development services to underserved communities, fostering economic empowerment and sustainable community growth.')
    
    # Background Summary
    p = doc.add_paragraph()
    p.add_run('Background Summary: ').bold = True
    p.add_run('Freedom Equity Inc has been serving the Cleveland, Ohio community for over 15 years. We specialize in affordable housing development, first-time homebuyer assistance, and small business incubation programs. Our organization has successfully developed 150+ affordable housing units and supported 200+ small businesses since our founding.')
    
    # Financial Information
    doc.add_heading('Financial Information', level=1)
    
    p = doc.add_paragraph()
    p.add_run('Annual Budget: ').bold = True
    p.add_run('$850,000')
    
    p = doc.add_paragraph()
    p.add_run('Primary Revenue Sources: ').bold = True
    p.add_run('Federal and state grants (60%), private foundation funding (25%), individual donations (15%)')
    
    # Contact Information
    doc.add_heading('Contact Information', level=1)
    
    p = doc.add_paragraph()
    p.add_run('Executive Director: ').bold = True
    p.add_run('Sarah Johnson')
    
    p = doc.add_paragraph()
    p.add_run('Email: ').bold = True
    p.add_run('sarah.johnson@freedomequity.org')
    
    p = doc.add_paragraph()
    p.add_run('Phone: ').bold = True
    p.add_run('216-555-0123')
    
    p = doc.add_paragraph()
    p.add_run('Address: ').bold = True
    p.add_run('1234 Community Drive, Cleveland, OH 44115')
    
    p = doc.add_paragraph()
    p.add_run('City: ').bold = True
    p.add_run('Cleveland, Ohio')
    
    # Programs and Services
    doc.add_heading('Programs and Services', level=1)
    
    doc.add_paragraph('• Affordable Housing Development Program')
    doc.add_paragraph('• First-Time Homebuyer Education and Assistance')
    doc.add_paragraph('• Small Business Incubator and Mentorship')
    doc.add_paragraph('• Financial Literacy Training')
    doc.add_paragraph('• Community Economic Development Initiatives')
    
    # Target Population
    doc.add_heading('Target Population', level=1)
    doc.add_paragraph('Low to moderate-income families and individuals in Cleveland and surrounding communities, with particular focus on minority-owned businesses and first-generation homebuyers.')
    
    # Save document
    filepath = Path('sample_organization_info.docx')
    doc.save(filepath)
    return filepath

def create_funder_document():
    """Create a realistic funder document"""
    doc = Document()
    
    doc.add_heading('KeyBank Foundation Bicentennial Grant Program', 0)
    
    doc.add_heading('Grant Program Overview', level=1)
    
    p = doc.add_paragraph()
    p.add_run('Foundation Name: ').bold = True
    p.add_run('KeyBank Foundation')
    
    p = doc.add_paragraph()
    p.add_run('Grant Program: ').bold = True
    p.add_run('Bicentennial Community Development Grant')
    
    p = doc.add_paragraph()
    p.add_run('Request Limit: ').bold = True
    p.add_run('$25,000')
    
    p = doc.add_paragraph()
    p.add_run('Grant Type: ').bold = True
    p.add_run('Unrestricted Operating Support')
    
    # Funding Priorities
    doc.add_heading('Funding Priorities', level=1)
    doc.add_paragraph('The KeyBank Foundation Bicentennial Grant Program supports organizations working in the following areas:')
    
    doc.add_paragraph('• Affordable housing development and homeownership programs')
    doc.add_paragraph('• Small business and entrepreneurship support initiatives')
    doc.add_paragraph('• Community economic development projects')
    doc.add_paragraph('• Financial inclusion and literacy programs')
    doc.add_paragraph('• Workforce development and job training')
    doc.add_paragraph('• Community revitalization efforts')
    
    # Eligibility Criteria
    doc.add_heading('Eligibility Criteria', level=1)
    doc.add_paragraph('To be eligible for funding, organizations must meet the following requirements:')
    
    doc.add_paragraph('• 501(c)(3) non-profit organizations in good standing')
    doc.add_paragraph('• Organizations located within KeyBank\'s footprint communities')
    doc.add_paragraph('• Demonstrated track record of community impact')
    doc.add_paragraph('• Organizations with annual budgets between $100,000 and $2,000,000')
    doc.add_paragraph('• Focus on serving low to moderate-income populations')
    
    # Grant Requirements
    doc.add_heading('Grant Requirements', level=1)
    doc.add_paragraph('Successful grantees will be required to:')
    
    doc.add_paragraph('• Submit quarterly progress reports')
    doc.add_paragraph('• Participate in capacity building opportunities')
    doc.add_paragraph('• Provide final impact report within 60 days of grant period end')
    doc.add_paragraph('• Acknowledge KeyBank Foundation support in all related materials')
    
    # Application Deadline
    doc.add_heading('Important Dates', level=1)
    p = doc.add_paragraph()
    p.add_run('Application Deadline: ').bold = True
    p.add_run('December 31, 2024')
    
    p = doc.add_paragraph()
    p.add_run('Award Notification: ').bold = True
    p.add_run('February 28, 2025')
    
    p = doc.add_paragraph()
    p.add_run('Grant Period: ').bold = True
    p.add_run('March 1, 2025 - February 28, 2026')
    
    # Save document
    filepath = Path('sample_funder_info.docx')
    doc.save(filepath)
    return filepath

def create_sample_documents():
    """Create both sample documents"""
    print("📄 Creating realistic sample documents for demonstration...")
    
    try:
        org_file = create_organization_document()
        print(f"✅ Created organization document: {org_file}")
        
        funder_file = create_funder_document()
        print(f"✅ Created funder document: {funder_file}")
        
        print("\n🎯 Sample documents ready for grant building demonstration!")
        return True
        
    except Exception as e:
        print(f"❌ Error creating sample documents: {e}")
        return False

if __name__ == "__main__":
    create_sample_documents() 